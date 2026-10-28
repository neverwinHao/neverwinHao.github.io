#!/usr/bin/env python3
"""Build the native four-slide PPTX with the cover and three embedded demo clips.

Run after record_demo.py. Requires python-pptx and ffmpeg. The resulting
DPIAgent-Demo.pptx is written one directory above the site checkout.
"""
from pathlib import Path
import json, subprocess, zipfile, xml.etree.ElementTree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / 'neverwinHao.github.io/DPIAgent/assets/demo'
WORK = Path('/tmp/dpi-preview/pptx-export')
WORK.mkdir(exist_ok=True)
OUT = ROOT / 'DPIAgent-Demo.pptx'
meta = json.loads((ASSETS / 'recording.json').read_text())
assert len(meta['scenes']) == 4 and meta['playback_rate'] == 3
prs = Presentation(ASSETS / 'dpiagent-cover.pptx')
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = 'DPIAgent'
prs.core_properties.subject = 'Reproduction test generation'
prs.core_properties.author = 'DPIAgent'

def slide(seconds):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    transition = OxmlElement('p:transition')
    transition.set('spd', 'slow')
    transition.set('advTm', str(int(seconds * 1000)))
    transition.set('advClick', '1')
    transition.append(OxmlElement('p:fade'))
    s._element.insert_element_before(transition, 'p:timing', 'p:extLst')
    return s

cover = prs.slides[0]
for old in cover._element.xpath('./p:transition'):
    cover._element.remove(old)
transition = OxmlElement('p:transition')
transition.set('spd','slow');transition.set('advTm',str(round(meta['scenes'][0]['duration']*1000)));transition.set('advClick','1')
transition.append(OxmlElement('p:fade'))
cover._element.insert_element_before(transition,'p:timing','p:extLst')

for scene in meta['scenes']:
    if scene['kind'] != 'demo':
        continue
    task = scene['task']
    record = next(r for r in meta['recordings'] if r['id'] == task)
    clip, poster = WORK / f'{task}.mp4', WORK / f'{task}.png'
    subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-y','-ss',str(scene['start']),'-i',str(ASSETS/'dpiagent-demo.mp4'),'-t',str(scene['duration']),'-an','-c:v','libx264','-threads','4','-preset','fast','-crf','18','-pix_fmt','yuv420p','-movflags','+faststart',str(clip)],check=True)
    subprocess.run(['ffmpeg','-hide_banner','-loglevel','error','-y','-ss','1','-i',str(clip),'-frames:v','1',str(poster)],check=True)
    s = slide(scene['duration'])
    s.shapes.add_movie(str(clip), 0, 0, prs.slide_width, prs.slide_height, poster_frame_image=str(poster), mime_type='video/mp4')
    for cond in s._element.xpath('.//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond'):
        cond.set('delay','0')
    url = f"https://neverwinhao.github.io/DPIAgent/demo/?task={task}&config={record['config']}"
    s.notes_slide.notes_text_frame.text = (
        f"{task}\n{url}\nEmbedded {scene['duration']}-second video, automatic playback in Slide Show. "
        f"Complete original trajectory at {meta['playback_rate']}x speed; every action and observation retained in order. "
        'Final FAILED/PASSED excerpts come from separate recorded reference-fix evaluation of the same submitted test patch.\n'
        f"Original trajectory SHA-256: {record['source_sha256']}\n" + json.dumps(record['evaluation'], indent=2)
    )
prs.save(OUT)
with zipfile.ZipFile(OUT) as z:
    videos = [n for n in z.namelist() if n.startswith('ppt/media/') and n.endswith('.mp4')]
    assert len(videos) == 3 and len(prs.slides) == 4
    for n in videos:
        assert len(z.read(n)) > 10000
    for n in z.namelist():
        if n.endswith('.rels'):
            for rel in ET.fromstring(z.read(n)):
                if rel.attrib.get('Type','').endswith(('/video','/media')):
                    assert rel.attrib.get('TargetMode') != 'External'
    for i in range(1,5):
        xml = ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
        ns = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
        assert xml.find('p:transition/p:fade', ns) is not None
print('Saved', OUT, 'slides', len(prs.slides), 'embedded videos', len(videos), 'bytes', OUT.stat().st_size, flush=True)
